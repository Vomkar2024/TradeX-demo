"""
TradeX Markdown to Word (.docx) Documentation Converter
Converts all technical specification markdown files in docs/ into
professionally styled, executive-grade Microsoft Word (.docx) documents.
"""

import os
import re
import sys
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

NAVY_COLOR = RGBColor(30, 58, 138)       # #1E3A8A
TEAL_COLOR = RGBColor(15, 118, 110)      # #0F766E
SLATE_COLOR = RGBColor(51, 65, 85)       # #334155
GRAY_TEXT = RGBColor(100, 116, 139)      # #64748B
CODE_TEXT = RGBColor(30, 41, 59)         # #1E293B

def set_cell_background(cell, fill_hex):
    """Sets cell background fill color using XML shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Sets cell internal padding/margins in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border_left_accent(cell, color_hex="2563EB", size_eighths="36"):
    """Adds a bold left accent border to a callout cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="{size_eighths}" w:space="0" w:color="{color_hex}"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_table_borders(table, border_color="CBD5E1"):
    """Sets clean subtle borders for standard data tables."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def parse_inline_formatting(paragraph, text):
    """
    Parses bold (**text**), italic (*text*), code (`text`), and links ([text](url))
    and adds formatted runs to paragraph.
    """
    # Pattern to match bold, italic, code, or links
    pattern = re.compile(
        r'(\*\*(.+?)\*\*)|'       # Group 1 & 2: Bold
        r'(\*(.+?)\*)|'           # Group 3 & 4: Italic
        r'(`(.+?)`)|'             # Group 5 & 6: Code
        r'(\[(.+?)\]\((.*?)\))'    # Group 7 & 8 & 9: Link
    )

    last_idx = 0
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            # Regular text
            run = paragraph.add_run(text[last_idx:start])
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)

        if match.group(1):  # Bold
            run = paragraph.add_run(match.group(2))
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.bold = True
        elif match.group(3):  # Italic
            run = paragraph.add_run(match.group(4))
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.italic = True
        elif match.group(5):  # Code
            run = paragraph.add_run(match.group(6))
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(185, 28, 28)  # Deep red/maroon for code
        elif match.group(7):  # Link
            link_text = match.group(8)
            run = paragraph.add_run(link_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(37, 99, 235)  # Blue link
            run.underline = True

        last_idx = end

    if last_idx < len(text):
        run = paragraph.add_run(text[last_idx:])
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)

def add_styled_heading(doc, text, level):
    """Adds a custom styled heading."""
    clean_text = re.sub(r'^[#\s]+', '', text).strip()
    # Remove emoji markers if desired or keep them
    h = doc.add_heading(level=min(level, 3))
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(4)

    run = h.add_run(clean_text)
    if level == 1:
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(16)
        run.font.color.rgb = NAVY_COLOR
        run.bold = True
    elif level == 2:
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(13)
        run.font.color.rgb = TEAL_COLOR
        run.bold = True
    else:
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(11)
        run.font.color.rgb = SLATE_COLOR
        run.bold = True
    return h

def add_code_block(doc, code_lines, language=""):
    """Adds a code block formatted inside a styled single-cell table."""
    code_text = "\n".join(code_lines)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_border_left_accent(cell, color_hex="0EA5E9", size_eighths="24")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    if language:
        lang_run = p.add_run(f"// Language: {language}\n")
        lang_run.font.name = 'Consolas'
        lang_run.font.size = Pt(8.5)
        lang_run.font.color.rgb = GRAY_TEXT
        lang_run.italic = True

    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.0)
    run.font.color.rgb = CODE_TEXT

    # Add small spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_callout_box(doc, lines, callout_type="NOTE"):
    """Adds an executive callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)

    bg_color = "F0FDF4" if "TIP" in callout_type or "SUCCESS" in callout_type else "EFF6FF"
    border_color = "16A34A" if "TIP" in callout_type or "SUCCESS" in callout_type else "2563EB"

    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_border_left_accent(cell, color_hex=border_color, size_eighths="32")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)

    title_run = p.add_run(f"📌 {callout_type}: ")
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(10)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)

    first = True
    for line in lines:
        if not first:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
        parse_inline_formatting(p, line)
        first = False

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def convert_markdown_to_docx(md_path, docx_path, doc_title="TradeX Documentation"):
    """Parses a markdown file and writes a styled .docx file."""
    print(f"Converting: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()

    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.different_first_page_header_footer = True

        # Header
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = header_p.add_run(f"TradeX Technical Specification  |  {doc_title}")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = GRAY_TEXT

        # Footer
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = footer_p.add_run("Confidential  •  Institutional Equity Simulation Desk")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = GRAY_TEXT

    # Determine Title from first H1
    h1_title = doc_title
    for line in lines:
        if line.startswith("# "):
            h1_title = line.replace("# ", "").strip()
            break

    # Add Document Cover / Title Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run(h1_title)
    t_run.font.name = 'Segoe UI Semibold'
    t_run.font.size = Pt(22)
    t_run.font.color.rgb = NAVY_COLOR
    t_run.bold = True

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    s_run = sub_p.add_run("TradeX Institutional Simulation Platform  •  Technical Specification Series")
    s_run.font.name = 'Calibri'
    s_run.font.size = Pt(11)
    s_run.font.color.rgb = TEAL_COLOR
    s_run.bold = True

    # Divider bar
    div_table = doc.add_table(rows=1, cols=1)
    div_cell = div_table.cell(0, 0)
    div_cell.width = Inches(6.9)
    set_cell_background(div_cell, "1E3A8A")
    set_cell_margins(div_cell, top=20, bottom=20, left=0, right=0)
    div_p = div_cell.paragraphs[0]
    div_p.paragraph_format.space_before = Pt(0)
    div_p.paragraph_format.space_after = Pt(0)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)

    i = 0
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    md_dir = os.path.dirname(os.path.abspath(md_path))
    workspace_root = os.path.abspath(os.path.join(md_dir, ".."))

    while i < len(lines):
        line = lines[i]

        # Handle Code Block fences
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line.strip().replace("```", "").strip()
                code_block_lines = []
            else:
                in_code_block = False
                add_code_block(doc, code_block_lines, code_block_lang)
                code_block_lines = []
                code_block_lang = ""
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Skip shields / badges
        if line.strip().startswith("[![") or ("img.shields.io" in line and "[!" in line):
            i += 1
            continue

        # Skip main document H1 (already in cover title)
        if line.startswith("# ") and i < 5:
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ["---", "***", "___"]:
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(4)
            sp.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Headings
        if line.startswith("#### "):
            add_styled_heading(doc, line[5:], level=3)
            i += 1
            continue
        elif line.startswith("### "):
            add_styled_heading(doc, line[4:], level=2)
            i += 1
            continue
        elif line.startswith("## "):
            add_styled_heading(doc, line[3:], level=1)
            i += 1
            continue
        elif line.startswith("# "):
            add_styled_heading(doc, line[2:], level=1)
            i += 1
            continue

        # Markdown Images: ![alt](path)
        img_match = re.search(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            alt_text = img_match.group(1)
            img_rel_path = img_match.group(2)
            # Resolve image path
            possible_paths = [
                os.path.normpath(os.path.join(md_dir, img_rel_path)),
                os.path.normpath(os.path.join(workspace_root, img_rel_path)),
                os.path.normpath(os.path.join(workspace_root, "assets", os.path.basename(img_rel_path)))
            ]
            img_found = None
            for p in possible_paths:
                if os.path.exists(p) and os.path.isfile(p):
                    img_found = p
                    break

            if img_found:
                try:
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(8)
                    img_p.paragraph_format.space_after = Pt(2)
                    doc.add_picture(img_found, width=Inches(6.0))

                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_before = Pt(2)
                    cap_p.paragraph_format.space_after = Pt(8)
                    c_run = cap_p.add_run(f"Figure: {alt_text or os.path.basename(img_found)}")
                    c_run.font.name = 'Calibri'
                    c_run.font.size = Pt(9.0)
                    c_run.italic = True
                    c_run.font.color.rgb = GRAY_TEXT
                except Exception as ex:
                    print(f"Warning: Could not embed image {img_found}: {ex}")
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(f"[Diagram: {alt_text}]")
                r.font.name = 'Calibri'
                r.italic = True
                r.font.color.rgb = GRAY_TEXT
            i += 1
            continue

        # Callouts / Blockquotes (> ...)
        if line.startswith(">"):
            callout_lines = []
            callout_type = "NOTE"
            while i < len(lines) and lines[i].startswith(">"):
                raw_c = lines[i].lstrip(">").strip()
                if raw_c.startswith("[!"):
                    m = re.match(r'\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]', raw_c)
                    if m:
                        callout_type = m.group(1)
                        raw_c = raw_c[m.end():].strip()
                if raw_c:
                    callout_lines.append(raw_c)
                i += 1
            if callout_lines:
                add_callout_box(doc, callout_lines, callout_type=callout_type)
            continue

        # Markdown Tables: lines containing '|'
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                # Parse table
                rows_data = []
                for tl in table_lines:
                    # Check if separator row like | :--- | :--- |
                    if re.match(r'^\|[\s:-|]+\|$', tl):
                        continue
                    cells = [c.strip() for c in tl.strip('|').split('|')]
                    rows_data.append(cells)

                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl)

                    for r_idx, row in enumerate(rows_data):
                        is_header = (r_idx == 0)
                        for c_idx in range(num_cols):
                            cell_val = row[c_idx] if c_idx < len(row) else ""
                            cell = tbl.cell(r_idx, c_idx)
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

                            if is_header:
                                set_cell_background(cell, "1E3A8A")
                            elif r_idx % 2 == 1:
                                set_cell_background(cell, "F8FAFC")
                            else:
                                set_cell_background(cell, "FFFFFF")

                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            p.paragraph_format.line_spacing = 1.1

                            if is_header:
                                run = p.add_run(cell_val)
                                run.font.name = 'Segoe UI Semibold'
                                run.font.size = Pt(9.5)
                                run.font.color.rgb = RGBColor(255, 255, 255)
                                run.bold = True
                            else:
                                parse_inline_formatting(p, cell_val)

                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(6)
            continue

        # Bullet List Items
        bullet_match = re.match(r'^(\s*)[-*+]\s+(.*)', line)
        if bullet_match:
            indent = len(bullet_match.group(1))
            item_text = bullet_match.group(2)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if indent >= 2:
                p.paragraph_format.left_indent = Inches(0.4)
            parse_inline_formatting(p, item_text)
            i += 1
            continue

        # Numbered List Items
        num_match = re.match(r'^(\s*)\d+\.\s+(.*)', line)
        if num_match:
            indent = len(num_match.group(1))
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if indent >= 2:
                p.paragraph_format.left_indent = Inches(0.4)
            parse_inline_formatting(p, item_text)
            i += 1
            continue

        # Normal Paragraph
        clean_line = line.strip()
        if clean_line:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, clean_line)

        i += 1

    doc.save(docx_path)
    file_size = os.path.getsize(docx_path)
    print(f"[OK] Saved {docx_path} ({file_size:,} bytes)")
    return file_size

def convert_all_docs():
    """Converts all 5 specification documents in docs/"""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    docs_dir = os.path.abspath(os.path.join(script_dir, "..", "docs"))

    specs = [
        ("ARCHITECTURE.md", "ARCHITECTURE.docx", "TradeX System Architecture Specification"),
        ("ER_DIAGRAMS.md", "ER_DIAGRAMS.docx", "TradeX Entity Relationship Diagrams & Schemas"),
        ("FEATURES.md", "FEATURES.docx", "TradeX Core Platform Features Specification"),
        ("MVC_PATTERN.md", "MVC_PATTERN.docx", "TradeX MVC Architecture & Data Flow Guide"),
        ("USER_FLOW.md", "USER_FLOW.docx", "TradeX End-to-End User Flow Specification")
    ]

    results = []
    print("\n" + "="*60)
    print("TradeX Technical Documentation -> Word (.docx) Converter")
    print("="*60)

    for md_file, docx_file, title in specs:
        md_full = os.path.join(docs_dir, md_file)
        docx_full = os.path.join(docs_dir, docx_file)
        if os.path.exists(md_full):
            sz = convert_markdown_to_docx(md_full, docx_full, title)
            results.append((docx_file, sz, "SUCCESS"))
        else:
            print(f"Error: {md_full} not found!")
            results.append((docx_file, 0, "FILE_NOT_FOUND"))

    print("\n" + "="*60)
    print("Conversion Summary:")
    print("="*60)
    for fname, size, status in results:
        print(f" • {fname:<22} : {size:>8,} bytes [{status}]")
    print("="*60 + "\n")

if __name__ == "__main__":
    convert_all_docs()
